import json
import re
from pathlib import Path
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


OUTPUT_JSON_DIR = Path("output/json")
OUTPUT_REPORTS_DIR = Path("output/reports")

ORDERED_HEADERS = [
    "Cliente",
    "Código JS",
    "Referência Cliente",
    "Responsável Atual",
    "Status",
    "Modal",
    "Tipo Transporte",
    "Regime",
    "Embarcação/Plataforma",
    "Incoterm",
    "ETA",
    "Data de Embarque",
    "Data de Chegada",
    "Data de Registro da DI",
    "Numero da DI",
    "Data da CI",
    "Canal",
    "House AWB/BL",
    "Master AWB/BL",
    "Pack List",
    "Commercial Invoice",
    "Proforma Invoice",
    "País de Origem",
    "Local de Embarque",
    "URF Entrada no País",
    "Peso Líquido",
    "Peso Bruto",
    "Volume",
    "Valor Declarado",
    "Moeda",
    "Valor da Taxa",
    "Frete Collect",
    "Frete Prepaid",
    "Frete Nacional",
    "Moeda Frete",
    "Recinto Aduaneiro",
    "Armazém",
]


def latest_consolidated_json() -> Path:
    files = sorted(OUTPUT_JSON_DIR.glob("followup_consolidado_*.json"))
    if not files:
        raise FileNotFoundError(
            "Nenhum followup_consolidado_*.json encontrado em output/json. "
            "Rode primeiro o main.py com os processos."
        )
    return files[-1]


def clean_simple_value(value):
    if value is None:
        return ""
    text = str(value).replace("\xa0", " ").strip()
    text = re.sub(r"[ \t]+", " ", text)
    return text


def first_meaningful_line(value):
    text = clean_simple_value(value)
    if not text:
        return ""

    stop_words = {
        "referência", "responsável", "status", "vinculado",
        "dados gerais", "documentos", "informações adicionais",
        "ocorrência", "demonstrativo de despesas"
    }

    for raw_line in str(value).splitlines():
        line = raw_line.strip(" .:\t\r\n")
        if not line:
            continue

        lowered = line.lower()
        if lowered in stop_words:
            break

        if any(lowered.startswith(sw) for sw in stop_words):
            break

        if line not in {"-", ":"}:
            return line

    return text


def normalize_modal(tipo_processo, tipo_transporte):
    joined = f"{clean_simple_value(tipo_processo)} {clean_simple_value(tipo_transporte)}".upper()

    if "AERE" in joined:
        return "AÉREO"
    if "MARIT" in joined:
        return "MARÍTIMO"
    if "RODOV" in joined:
        return "RODOVIÁRIO"
    return clean_simple_value(tipo_processo) or clean_simple_value(tipo_transporte)


def get_section_tables(section):
    if not isinstance(section, dict):
        return []
    return section.get("tables", []) or []


def extract_doc_rows(section):
    """
    Converte tabelas de Documentos / Informações Adicionais em lista de dicionários.
    Espera linhas no formato:
    [Tipo Documento, Nr.Documento, Data Entrega, Data Validade, ...]
    """
    results = []

    for table in get_section_tables(section):
        rows = table.get("rows", [])
        if len(rows) < 2:
            continue

        header = [clean_simple_value(x).lower() for x in rows[0]]
        joined = " | ".join(header)

        if "tipo documento" not in joined or "nr.documento" not in joined:
            continue

        for row in rows[1:]:
            if not row:
                continue

            first = clean_simple_value(row[0]).lower()
            if first.startswith("total"):
                continue

            results.append({
                "tipo_documento": clean_simple_value(row[0]) if len(row) > 0 else "",
                "nr_documento": clean_simple_value(row[1]) if len(row) > 1 else "",
                "data_entrega": clean_simple_value(row[2]) if len(row) > 2 else "",
                "data_validade": clean_simple_value(row[3]) if len(row) > 3 else "",
            })

    return results


def find_doc(rows, *wanted_types):
    wanted = {w.upper() for w in wanted_types}
    for row in rows:
        if clean_simple_value(row.get("tipo_documento")).upper() in wanted:
            return row
    return None


def first_non_empty(*values):
    for value in values:
        value = clean_simple_value(value)
        if value:
            return value
    return ""


def extract_from_preview(preview, label):
    """
    Procura coisas do tipo:
    LABEL=>valor
    """
    text = clean_simple_value(preview)
    if not text:
        return ""

    pattern = rf"{re.escape(label)}\s*=>\s*([^\r\n]+)"
    match = re.search(pattern, text, flags=re.IGNORECASE)
    return clean_simple_value(match.group(1)) if match else ""


def parse_data_chegada_from_preview(preview):
    text = clean_simple_value(preview)
    match = re.search(r"Data Chegada Carga:\s*([0-9]{2}/[0-9]{2}/[0-9]{4})", text, flags=re.IGNORECASE)
    return clean_simple_value(match.group(1)) if match else ""


def build_row(process_data):
    resumo = process_data.get("resumo", {}) or {}
    dados = process_data.get("dados_gerais", {}) or {}
    documentos = process_data.get("documentos", {}) or {}
    adicionais = process_data.get("documentos_adicionais", {}) or {}
    ocorrencia = process_data.get("ocorrencia", {}) or {}

    doc_rows = extract_doc_rows(documentos)
    add_rows = extract_doc_rows(adicionais)
    rows_all = doc_rows + add_rows

    doc_di = find_doc(rows_all, "DI")
    doc_ci = find_doc(rows_all, "CI", "CI (COMPROV. IMPORT)")
    doc_house = find_doc(rows_all, "HAWB/HBL")
    doc_master = find_doc(rows_all, "AWB/MAWB", "MASTER BL")
    doc_commercial = find_doc(rows_all, "INVOICE - COMMERCIAL")
    doc_proforma = find_doc(rows_all, "INVOICE - PROFORMA", "PROFORMA")
    doc_ref_cliente = find_doc(rows_all, "REFERENCIA CLIENTE")
    doc_pais_origem = find_doc(rows_all, "PAÍS DE ORIGEM")
    doc_peso_bruto = find_doc(rows_all, "PESO BRUTO")
    doc_volumes = find_doc(rows_all, "VOLUMES")
    doc_data_registro_di = find_doc(rows_all, "DATA REGISTRO DI")
    doc_canal = find_doc(rows_all, "CANAL")
    doc_eta = find_doc(rows_all, "PREVISÃO DE CHEGADA", "ETA")
    doc_pack_list = find_doc(rows_all, "ROMANEIO CARGA / PACKING LIST", "PACK LIST")

    occ_preview = ocorrencia.get("preview", "")
    dados_preview = dados.get("preview", "")

    cliente = first_meaningful_line(resumo.get("cliente"))
    referencia_resumo = first_meaningful_line(resumo.get("referencia"))

    referencia_cliente = first_non_empty(
        doc_ref_cliente["nr_documento"] if doc_ref_cliente else "",
        extract_from_preview(occ_preview, "REFERENCIA CLIENTE"),
        referencia_resumo,
    )

    numero_di = first_non_empty(
        dados.get("nr_di"),
        doc_di["nr_documento"] if doc_di else "",
    )

    data_registro_di = first_non_empty(
        dados.get("dt_registro_di"),
        doc_data_registro_di["nr_documento"] if doc_data_registro_di else "",
        doc_data_registro_di["data_entrega"] if doc_data_registro_di else "",
    )

    data_ci = first_non_empty(
        doc_ci["data_entrega"] if doc_ci else "",
        extract_from_preview(occ_preview, "DATA DA CI"),
    )

    canal = first_non_empty(
        doc_canal["nr_documento"] if doc_canal else "",
        extract_from_preview(occ_preview, "CANAL"),
    )

    eta = first_non_empty(
        doc_eta["nr_documento"] if doc_eta else "",
        extract_from_preview(occ_preview, "PREVISÃO DE CHEGADA"),
        extract_from_preview(occ_preview, "ETA"),
    )

    data_chegada = first_non_empty(
        dados.get("data_chegada_carga"),
        parse_data_chegada_from_preview(dados_preview),
    )

    house_awb_bl = first_non_empty(
        doc_house["nr_documento"] if doc_house else "",
    )

    master_awb_bl = first_non_empty(
        dados.get("identificacao_master"),
        doc_master["nr_documento"] if doc_master else "",
    )

    commercial_invoice = first_non_empty(
        doc_commercial["nr_documento"] if doc_commercial else "",
    )

    proforma_invoice = first_non_empty(
        doc_proforma["nr_documento"] if doc_proforma else "",
    )

    pais_origem = first_non_empty(
        doc_pais_origem["nr_documento"] if doc_pais_origem else "",
        dados.get("pais_procedencia_carga"),
    )

    peso_bruto = first_non_empty(
        doc_peso_bruto["nr_documento"] if doc_peso_bruto else "",
        dados.get("peso_bruto"),
    )

    volume = first_non_empty(
        doc_volumes["nr_documento"] if doc_volumes else "",
    )

    pack_list = first_non_empty(
        dados.get("pack_list"),
        doc_pack_list["nr_documento"] if doc_pack_list else "",
    )

    return {
        "Cliente": cliente,
        "Código JS": process_data.get("process_id", ""),
        "Referência Cliente": referencia_cliente,
        "Responsável Atual": clean_simple_value(resumo.get("responsavel_atual")),
        "Status": clean_simple_value(resumo.get("status_processo")),
        "Modal": normalize_modal(dados.get("tipo_processo"), dados.get("tipo_transporte")),
        "Tipo Transporte": clean_simple_value(dados.get("tipo_transporte")),
        "Regime": clean_simple_value(dados.get("tipo_declaracao")),
        "Embarcação/Plataforma": clean_simple_value(dados.get("embarcacao_plataforma")),
        "Incoterm": clean_simple_value(dados.get("incoterm")),
        "ETA": eta,
        "Data de Embarque": clean_simple_value(dados.get("data_embarque")),
        "Data de Chegada": data_chegada,
        "Data de Registro da DI": data_registro_di,
        "Numero da DI": numero_di,
        "Data da CI": data_ci,
        "Canal": canal,
        "House AWB/BL": house_awb_bl,
        "Master AWB/BL": master_awb_bl,
        "Pack List": pack_list,
        "Commercial Invoice": commercial_invoice,
        "Proforma Invoice": proforma_invoice,
        "País de Origem": pais_origem,
        "Local de Embarque": clean_simple_value(dados.get("local_embarque")),
        "URF Entrada no País": clean_simple_value(dados.get("urf_entrada_pais")),
        "Peso Líquido": clean_simple_value(dados.get("peso_liquido")),
        "Peso Bruto": peso_bruto,
        "Volume": volume,
        "Valor Declarado": clean_simple_value(dados.get("valor_declarado")),
        "Moeda": clean_simple_value(dados.get("moeda")),
        "Valor da Taxa": clean_simple_value(dados.get("valor_taxa_cambio")),
        "Frete Collect": clean_simple_value(dados.get("frete_collect")),
        "Frete Prepaid": clean_simple_value(dados.get("frete_prepaid")),
        "Frete Nacional": clean_simple_value(dados.get("frete_nacional")),
        "Moeda Frete": clean_simple_value(dados.get("moeda_frete")),
        "Recinto Aduaneiro": extract_from_preview(dados_preview, "Recinto Aduaneiro"),
        "Armazém": extract_from_preview(dados_preview, "Armazém"),
    }


def autosize_columns(ws):
    for column_cells in ws.columns:
        length = 0
        col_letter = get_column_letter(column_cells[0].column)
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            length = max(length, len(value))
        ws.column_dimensions[col_letter].width = min(length + 2, 35)


def export_xlsx(rows, output_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Coordenação"

    headers = ORDERED_HEADERS
    ws.append(headers)

    for row in rows:
        ws.append([row.get(h, "") for h in headers])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    autosize_columns(ws)

    wb.save(output_path)


def main():
    consolidated_path = latest_consolidated_json()

    with open(consolidated_path, "r", encoding="utf-8") as f:
        data = json.load(f)

    processos = []
    for _, items in (data.get("clientes") or {}).items():
        processos.extend(items)

    rows = [build_row(p) for p in processos]

    OUTPUT_REPORTS_DIR.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_path = OUTPUT_REPORTS_DIR / f"followup_coordenadora_{timestamp}.xlsx"

    export_xlsx(rows, output_path)

    print(f"Planilha criada em: {output_path}")
    print(f"Total de linhas: {len(rows)}")


if __name__ == "__main__":
    main()